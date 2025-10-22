# api/convert.py
import json
import base64
from PyPDF2 import PdfReader
from docx import Document
from io import BytesIO

def handler(request):
    """
    PDF file (base64 encoded) leta hai aur use Word (DOCX) file mein convert karke 
    base64 string ki shakal mein wapis bhejta hai.
    """
    
    # 1. Incoming Request (JSON) ko Parse karna
    try:
        data = request.json
        if not data or 'pdf_base64' not in data:
            return json.dumps({"error": "No PDF data found"}), 400, {'Content-Type': 'application/json'}
        
        # Base64 string se binary PDF data decode karna
        pdf_data = base64.b64decode(data['pdf_base64'])
        pdf_file = BytesIO(pdf_data)

    except Exception as e:
        return json.dumps({"error": f"Invalid input format: {e}"}), 400, {'Content-Type': 'application/json'}

    # 2. PDF to DOCX Conversion Logic (Simple Text Extraction)
    try:
        pdf_reader = PdfReader(pdf_file)
        document = Document()
        
        # Har page se text nikalna
        for page in pdf_reader.pages:
            text = page.extract_text()
            if text:
                document.add_paragraph(text)
                # Har page ke baad ek line break
                document.add_paragraph('--- NEW PAGE ---') 
        
        # 3. DOCX file ko memory (BytesIO) mein save karna
        docx_file = BytesIO()
        document.save(docx_file)
        docx_file.seek(0)
        
        # 4. DOCX data ko base64 mein encode karna
        docx_base64 = base64.b64encode(docx_file.read()).decode('utf-8')
        
        # 5. Result wapis Android App ko bhejna
        return json.dumps({
            "status": "success",
            "docx_base64": docx_base64,
            "filename": "converted_document.docx"
        }), 200, {'Content-Type': 'application/json'}

    except Exception as e:
        # Koi bhi error ho to use handle karna
        return json.dumps({"error": f"Conversion failed: {e}"}), 500, {'Content-Type': 'application/json'}